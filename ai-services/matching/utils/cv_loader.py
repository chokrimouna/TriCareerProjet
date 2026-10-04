import os
import fitz  # PyMuPDF (PDF)
from config import CV_FOLDER


def _read_pdf(path):
    text = ""
    doc = fitz.open(path)
    for page in doc:
        text += page.get_text()
    doc.close()
    return text


def _read_docx(path):
    """Extrait le texte d'un CV Word .docx (paragraphes + tableaux)."""
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def load_cv(filename):
    """Charge le texte d'un CV depuis CV_FOLDER, PDF ou Word (.docx).

    `filename` peut être fourni AVEC extension (ex. "cv.docx") ou SANS (ex. "cv",
    compat historique où l'appelant retirait ".pdf"). Dans ce dernier cas on cherche
    .pdf puis .docx. Renvoie "" si introuvable ou format non lisible.
    """
    if not filename:
        return ""

    _base, ext = os.path.splitext(filename)
    if ext:
        candidates = [filename]
    else:
        candidates = [filename + ".pdf", filename + ".docx", filename + ".doc"]

    path = None
    for c in candidates:
        p = os.path.join(CV_FOLDER, c)
        if os.path.exists(p):
            path = p
            break
    if not path:
        return ""

    lower = path.lower()
    try:
        if lower.endswith(".pdf"):
            return _read_pdf(path)
        if lower.endswith(".docx"):
            return _read_docx(path)
        # .doc binaire ancien : non pris en charge par python-docx
        return ""
    except Exception:
        return ""
